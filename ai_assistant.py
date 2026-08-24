"""
Интеграция с DeepSeek API для раздела "ИИ ассистент" — прогноз баллов
подрядчика по гипотетическому будущему конкурсу на основе истории его
реальных заявок (crud.get_supplier_ai_profile) и конкурентной среды
выбранного региона (crud.list_suppliers_in_region).

DeepSeek API совместим с форматом OpenAI Chat Completions, поэтому вызывается
обычным POST через requests — без отдельной SDK-зависимости.
"""

import os
from io import BytesIO

import requests
import truststore
from docx import Document
from dotenv import dotenv_values

# На части рабочих компьютеров (антивирус/корпоративный TLS-прокси с
# собственным перехватывающим сертификатом — та же причина, по которой
# порталу госзакупок в scraper.py пришлось отключать проверку сертификата)
# запросы к внешним HTTPS-адресам падают с "certificate verify failed:
# unable to get local issuer certificate": Windows этому сертификату
# доверяет, а отдельный список доверенных сертификатов, встроенный в
# Python (certifi), — нет. truststore переключает проверку на системное
# хранилище сертификатов Windows вместо этого списка, поэтому проверка
# подлинности сертификата не отключается (в отличие от verify=False).
truststore.inject_into_ssl()

_env = dotenv_values(".env")
DEEPSEEK_API_KEY = _env.get("DEEPSEEK_API_KEY") or os.getenv("DEEPSEEK_API_KEY")
DEEPSEEK_API_URL = "https://api.deepseek.com/chat/completions"
DEEPSEEK_MODEL = "deepseek-chat"

SYSTEM_PROMPT = (
    "Ты — аналитик по государственным закупкам Казахстана. Прогнозируешь "
    "баллы подрядчика по будущим конкурсам на основе истории его реальных "
    "заявок. Отвечай на русском языке, структурированно, по делу, без "
    "лишних оговорок."
)


class DeepSeekNotConfigured(Exception):
    """DEEPSEEK_API_KEY не задан в .env."""


def _format_history(history: list) -> str:
    if not history:
        return "У подрядчика в базе ещё нет истории участия в конкурсах."

    lines = []
    for h in history:
        result = (
            "победитель" if h["is_winner"]
            else "2 место" if h["is_second_place"]
            else (h["status"] or "участие")
        )
        crit_str = "; ".join(
            f"{c['name']} — значение: {c['value']}, балл: {c['score']}" for c in h["criteria"]
        ) or "критерии не зафиксированы в протоколе"
        line = (
            f"- Конкурс №{h['tender_no']} «{h['title'] or '—'}» "
            f"(категория: {', '.join(h['categories']) or '—'}; "
            f"сумма конкурса: {h['total_sum'] if h['total_sum'] else '—'} тенге; "
            f"участников: {h['participants_count']}; результат: {result}"
        )
        if h["rejection_reason"]:
            line += f"; причина отклонения: {h['rejection_reason']}"
        line += f"). Критерии и баллы этой заявки: {crit_str}"
        lines.append(line)
    return "\n".join(lines)


def _format_competitors(competitors: list) -> str:
    if not competitors:
        return "В этом регионе других зарегистрированных подрядчиков в базе пока нет."
    return "\n".join(
        f"- {c['name']} (БИН {c['bin']}): участий в конкурсах — {c['bids_count']}, побед — {c['wins_count']}"
        for c in competitors
    )


def build_prompt(supplier: dict, region_name: str, region_number: int,
                  future_sum: float, competitors: list) -> str:
    return f"""На основе данных ниже спрогнозируй результат подрядчика в будущем конкурсе.

ПОДРЯДЧИК: {supplier['name']} (БИН {supplier['bin']}), регион регистрации: {supplier['region_name'] or 'не указан'}.

ИСТОРИЯ УЧАСТИЯ ПОДРЯДЧИКА В КОНКУРСАХ (все заявки, известные базе):
{_format_history(supplier['history'])}

БУДУЩИЙ КОНКУРС, ПО КОТОРОМУ НУЖЕН ПРОГНОЗ:
- Регион проведения: {region_name} (код {region_number})
- Предполагаемая сумма конкурса: {future_sum:,.0f} тенге

ДРУГИЕ ПОДРЯДЧИКИ ЭТОГО РЕГИОНА В БАЗЕ (потенциальные конкуренты):
{_format_competitors(competitors)}

Проанализируй, в каких категориях работ подрядчик обычно участвует, как часто побеждает,
какие баллы обычно набирает по каждому критерию оценки и как это соотносится с суммой
конкурса и числом участников. На основе этого дай:

1. Прогноз баллов по каждому критерию оценки, которые подрядчик, вероятно, заявит по
   этому будущему конкурсу — с кратким обоснованием на основе его истории.
2. Итоговый прогнозный балл и вероятность победы (низкая/средняя/высокая) с учётом
   конкурентной среды региона.
3. Краткие рекомендации подрядчику, что усилить в заявке, чтобы повысить шансы на победу.

Если истории заявок недостаточно для уверенного прогноза — прямо скажи об этом и дай
осторожную оценку на основе того, что есть."""


def generate_forecast(supplier: dict, region_name: str, region_number: int,
                       future_sum: float, competitors: list) -> str:
    """Собирает промпт из данных БД и запрашивает прогноз у DeepSeek.
    Бросает DeepSeekNotConfigured, если ключ ещё не добавлен в .env."""
    if not DEEPSEEK_API_KEY:
        raise DeepSeekNotConfigured(
            "Не задан DEEPSEEK_API_KEY в .env — добавьте ключ API DeepSeek "
            "(строка DEEPSEEK_API_KEY=... в файле .env) и перезапустите приложение."
        )

    prompt = build_prompt(supplier, region_name, region_number, future_sum, competitors)

    resp = requests.post(
        DEEPSEEK_API_URL,
        headers={
            "Authorization": f"Bearer {DEEPSEEK_API_KEY}",
            "Content-Type": "application/json",
        },
        json={
            "model": DEEPSEEK_MODEL,
            "messages": [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": prompt},
            ],
            "temperature": 0.3,
        },
        timeout=120,
    )
    resp.raise_for_status()
    return resp.json()["choices"][0]["message"]["content"]


def _add_bold_runs(paragraph, text: str):
    """Разбивает строку по **...** (markdown-жирный, которым DeepSeek обычно
    выделяет ключевые фразы) и оформляет такие куски жирным шрифтом."""
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        paragraph.add_run(part).bold = (i % 2 == 1)


def _add_markdown_line(doc: Document, line: str):
    """Достаточно простой разбор markdown-ответа модели (заголовки #, списки
    -/*, **жирный**) — не полноценный парсер, а ровно то, что нужно, чтобы
    отчёт в Word не выглядел одной стеной текста со звёздочками и решётками."""
    stripped = line.strip()
    if not stripped:
        return

    heading_level = 0
    while stripped.startswith("#"):
        heading_level += 1
        stripped = stripped[1:]
    stripped = stripped.strip()
    if heading_level:
        doc.add_heading(stripped, level=min(heading_level, 4))
        return

    bullet = stripped.startswith(("- ", "* "))
    if bullet:
        stripped = stripped[2:]

    paragraph = doc.add_paragraph(style="List Bullet" if bullet else None)
    _add_bold_runs(paragraph, stripped)


def build_docx_report(supplier_name: str, region_name: str, future_sum: float,
                       created_at_str: str, report_text: str) -> bytes:
    """Собирает .docx с шапкой (подрядчик/регион/сумма/дата) и текстом
    прогноза — для скачивания и передачи отчёта тендерному отделу."""
    doc = Document()
    doc.add_heading("Прогноз ИИ-ассистента — МЭКС analytics", level=0)

    info = doc.add_paragraph()
    for label, value in (
        ("Подрядчик: ", supplier_name),
        ("Регион будущего конкурса: ", region_name),
        ("Предполагаемая сумма конкурса: ", f"{future_sum:,.0f} тенге"),
        ("Дата генерации отчёта: ", created_at_str),
    ):
        info.add_run(label).bold = True
        info.add_run(value + "\n")

    doc.add_paragraph()
    for line in report_text.splitlines():
        _add_markdown_line(doc, line)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
